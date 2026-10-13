"""Extract plain text from uploaded CVs (PDF, DOCX or TXT)."""
import io

from docx import Document
from pypdf import PdfReader

ALLOWED = {".pdf", ".docx", ".txt"}


class CVParseError(ValueError):
    pass


def extension(filename: str) -> str:
    name = (filename or "").lower()
    return name[name.rfind("."):] if "." in name else ""


def extract_text(filename: str, data: bytes) -> str:
    ext = extension(filename)
    if ext not in ALLOWED:
        raise CVParseError("Only PDF, DOCX or TXT files are supported.")
    try:
        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif ext == ".docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8", errors="ignore")
    except Exception as exc:  # corrupt or encrypted files
        raise CVParseError("We could not read this file. Try exporting it again, or paste the text instead.") from exc
    text = text.strip()
    if len(text) < 20:
        raise CVParseError("No readable text was found. Scanned PDFs need to be pasted as text.")
    return text
